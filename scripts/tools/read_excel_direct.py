from pathlib import Path

from openpyxl import load_workbook

SAMPLES = Path(__file__).resolve().parents[2] / "data" / "samples"
# Load workbook
wb = load_workbook(SAMPLES / "Data_SOAP_notes.xlsx")
sheet = wb.active

print("=" * 80)
print("SHEET INFO")
print("=" * 80)
print(f"Sheet name: {sheet.title}")
print(f"Total rows: {sheet.max_row}")
print(f"Total columns: {sheet.max_column}")
print("\n")

# Get headers (first row)
headers = []
for col in range(1, sheet.max_column + 1):
    headers.append(sheet.cell(1, col).value)
print(f"Headers: {headers}")
print("\n")

# Display first 3 examples
for row_num in range(2, min(5, sheet.max_row + 1)):
    print("=" * 80)
    print(f"EXAMPLE {row_num - 1}")
    print("=" * 80)
    
    for col_num, header in enumerate(headers, 1):
        value = sheet.cell(row_num, col_num).value
        if value:
            print(f"\n{header}:")
            print("-" * 40)
            print(value)
    print("\n")
