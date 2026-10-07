from pathlib import Path
import json

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "data" / "samples"

wb = load_workbook(SAMPLES / "Data_SOAP_notes.xlsx")
sheet = wb.active

# Save 3 examples to separate files for easier reading
for example_num in range(1, 4):
    row = example_num + 1  # +1 because row 1 is headers

    dialogue = sheet.cell(row, 1).value  # Column A
    soap_note = sheet.cell(row, 2).value  # Column B

    output = {
        "example_number": example_num,
        "dialogue": dialogue,
        "soap_note": soap_note,
    }

    filename = SAMPLES / f"soap_example_{example_num}.json"
    with open(filename, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    print(f"Saved {filename}")

print("\nDone! Saved 3 examples to JSON files.")
